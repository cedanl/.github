#!/usr/bin/env python3
"""Bestaande .pptx aanpassen zonder opnieuw te bouwen: lezen, tekst vervangen, logo wisselen, slides klonen, importeren, verplaatsen of verwijderen.

Voor decks die je niet zelf met build.mjs hebt gemaakt (bijvoorbeeld een voorstel dat al gedeeld is, of een
deck van een collega). Voor een nieuw deck gebruik je build.mjs en een deck-spec, geen edit.py.

Vereist: pip install python-pptx

Gebruik:
  python edit.py dump      in.pptx [--notes]
  python edit.py replace   in.pptx uit.pptx map.json [--lenient] [--slide N ...]
  python edit.py beeldmerk in.pptx uit.pptx
  python edit.py media     in.pptx uit.pptx ppt/media/image1.png nieuw.png [...paren] [--slide N ...]
  python edit.py kloon     in.pptx uit.pptx --slide N --naar M
  python edit.py importeer in.pptx uit.pptx bron.pptx --slide N --naar M
  python edit.py verplaats in.pptx uit.pptx --slide N --naar M
  python edit.py verwijder in.pptx uit.pptx --slide N [N ...]

Slidenummers beginnen bij 1. `--naar M` is de positie die de slide daarna heeft.
`importeer` zet slide N uit een ander bestand (bijvoorbeeld een met build.mjs gebouwde slide) op positie M.
`--slide` bij replace en media beperkt de wijziging tot die slides; nodig na `kloon`, want de kloon deelt
tekst en afbeeldingen met het origineel.
"""
import argparse
import copy
import io
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

for _stroom in (sys.stdout, sys.stderr):  # Windows-consoles gebruiken anders cp1252 en struikelen over tekens als →
    _stroom.reconfigure(encoding="utf-8")

# Logo's en beeldmerk staan in de projectroot van clidev-presentaties; draai edit.py vanuit die map.
ASSETS = Path("public/npuls/logos")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
R_EMBED = f"{R_NS}embed"


# ---------- lezen ----------

def tekstframes(slide, notes=False):
    for sh in slide.shapes:
        yield from frames_in(sh)
    if notes and slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        yield slide.notes_slide.notes_text_frame


def frames_in(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            yield from frames_in(sub)
    elif shape.has_text_frame:
        yield shape.text_frame
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def dump(inp, notes):
    prs = Presentation(inp)
    for n, slide in enumerate(prs.slides, 1):
        print(f"--- slide {n} ({slide.slide_layout.name})")
        for sh in slide.shapes:
            tekst = " | ".join(f.text.replace("\n", " / ") for f in frames_in(sh) if f.text.strip())
            soort = sh.shape_type
            if tekst:
                print(f"  [{sh.shape_id}] {sh.name}: {tekst}")
            elif soort == MSO_SHAPE_TYPE.PICTURE:
                print(f"  [{sh.shape_id}] {sh.name}: (afbeelding {str(slide.part.related_part(sh._element.blip_rId).partname).lstrip('/')})")
        if notes and slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            t = slide.notes_slide.notes_text_frame.text.strip()
            if t:
                print(f"  notities: {t[:200]}")


# ---------- tekst vervangen ----------

def vervang_in_alinea(para, oud, nieuw):
    """Vervangt `oud` in een alinea, ook als het over meerdere runs is verdeeld. Geeft het aantal vervangingen terug."""
    n = 0
    for r in para.runs:
        if oud in r.text:
            n += r.text.count(oud)
            r.text = r.text.replace(oud, nieuw)
    if n:
        return n
    volledig = "".join(r.text for r in para.runs)
    if oud in volledig and para.runs:  # opmaak van de eerste run blijft behouden
        n = volledig.count(oud)
        para.runs[0].text = volledig.replace(oud, nieuw)
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
    return n


def gekozen_slides(prs, nummers):
    if not nummers:
        return list(prs.slides)
    fout = [n for n in nummers if not 1 <= n <= len(prs.slides)]
    if fout:
        sys.exit(f"FOUT: slide {', '.join(map(str, fout))} bestaat niet (het deck heeft {len(prs.slides)} slides)")
    return [prs.slides[n - 1] for n in sorted(set(nummers))]


def replace(inp, uit, mapfile, lenient, nummers=None):
    mapping = json.loads(Path(mapfile).read_text(encoding="utf-8"))
    prs = Presentation(inp)
    telling = {oud: 0 for oud in mapping}
    for slide in gekozen_slides(prs, nummers):
        for tf in tekstframes(slide, notes=True):
            for para in tf.paragraphs:
                for oud, nieuw in mapping.items():
                    telling[oud] += vervang_in_alinea(para, oud, nieuw)
    ontbreekt = [o for o, n in telling.items() if n == 0]
    for oud, n in telling.items():
        print(f"{n:3d}x  {oud[:70]}")
    if ontbreekt and not lenient:
        sys.exit(f"FOUT: {len(ontbreekt)} zoekterm(en) niet gevonden (gebruik --lenient om dat toe te staan)")
    prs.save(uit)
    print("->", uit)


# ---------- beeldmerk ----------

def voorinstelling(shape):
    geom = shape._element.spPr.find(f"{A}prstGeom") if hasattr(shape._element, "spPr") else None
    return geom.get("prst") if geom is not None else None


def stipkleur(shape):
    try:
        rgb = str(shape.fill.fore_color.rgb)
    except Exception:
        return "zwart"
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return "wit" if 0.299 * r + 0.587 * g + 0.114 * b > 140 else "zwart"


def beeldmerk(inp, uit):
    """Vervangt een zelfgetekend of indicatief stippenmerk (>= 20 kleine cirkels) door het echte Npuls-beeldmerk."""
    max_stip = int(0.12 * 914400)
    prs = Presentation(inp)
    vervangen = 0
    for slide in prs.slides:
        stippen = [s for s in slide.shapes
                   if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and voorinstelling(s) == "ellipse"
                   and s.width <= max_stip and s.height <= max_stip]
        if len(stippen) < 20:
            continue
        links = min(s.left for s in stippen)
        boven = min(s.top for s in stippen)
        rechts = max(s.left + s.width for s in stippen)
        onder = max(s.top + s.height for s in stippen)
        zijde = max(rechts - links, onder - boven)
        cx, cy = (links + rechts) // 2, (boven + onder) // 2
        kleur = stipkleur(stippen[0])
        for s in stippen:
            s._element.getparent().remove(s._element)
        bestand = ASSETS / f"npuls-beeldmerk-{kleur}.png"
        if not bestand.exists():
            sys.exit(f"FOUT: {bestand} niet gevonden. Draai edit.py vanuit de root van clidev-presentaties.")
        pic = slide.shapes.add_picture(str(bestand), cx - zijde // 2, cy - zijde // 2, zijde, zijde)
        pic.name = "Npuls beeldmerk"
        vervangen += 1
    prs.save(uit)
    print(f"beeldmerk vervangen op {vervangen} slides -> {uit}")


# ---------- afbeeldingen ----------

def media(inp, uit, paren, nummers=None):
    """Vervangt de inhoud van een afbeelding in het bestand (bijvoorbeeld ppt/media/image1.png, zonder
    beginslash: Git Bash maakt van een beginslash een Windows-pad). Alle slides die die afbeelding gebruiken, krijgen
    de nieuwe. Met `nummers` krijgen alleen die slides een eigen kopie met de nieuwe afbeelding."""
    if len(paren) % 2:
        sys.exit("FOUT: geef paren van onderdeelnaam en nieuw bestand")
    prs = Presentation(inp)
    if nummers:
        gewenst = {paren[i].lstrip("/"): paren[i + 1] for i in range(0, len(paren), 2)}
        gevonden = set()
        for slide in gekozen_slides(prs, nummers):
            for blip in slide._element.iter(f"{A}blip"):
                rid = blip.get(R_EMBED)
                naam = str(slide.part.related_part(rid).partname).lstrip("/") if rid else None
                if naam in gewenst:
                    _, nieuw_rid = slide.part.get_or_add_image_part(gewenst[naam])
                    blip.set(R_EMBED, nieuw_rid)
                    gevonden.add(naam)
        ontbreekt = set(gewenst) - gevonden
        if ontbreekt:
            sys.exit(f"FOUT: niet gevonden op slide {', '.join(map(str, nummers))}: {', '.join(sorted(ontbreekt))}")
        prs.save(uit)
        print(f"{len(gevonden)} afbeelding(en) vervangen op slide {', '.join(map(str, nummers))} -> {uit}")
        return
    gewenst = {paren[i].lstrip("/"): Path(paren[i + 1]).read_bytes() for i in range(0, len(paren), 2)}
    gevonden = set()
    for part in prs.part.package.iter_parts():
        naam = str(part.partname).lstrip("/")
        if naam in gewenst:
            part._blob = gewenst[naam]
            gevonden.add(naam)
    ontbreekt = set(gewenst) - gevonden
    if ontbreekt:
        sys.exit(f"FOUT: niet gevonden in het bestand: {', '.join(sorted(ontbreekt))}")
    prs.save(uit)
    print(f"{len(gevonden)} afbeelding(en) vervangen -> {uit}")


# ---------- slides ----------

def verplaats_slide(prs, van, naar):
    lijst = prs.slides._sldIdLst
    ids = list(lijst)
    el = ids[van - 1]
    lijst.remove(el)
    lijst.insert(naar - 1, el)


def kloon(inp, uit, nummer, naar):
    prs = Presentation(inp)
    bron = prs.slides[nummer - 1]
    nieuw = prs.slides.add_slide(bron.slide_layout)
    for sh in list(nieuw.shapes):
        sh._element.getparent().remove(sh._element)

    def met_eigen_afbeeldingen(el):
        for blip in el.iter(f"{A}blip"):
            rid = blip.get(R_EMBED)
            if rid:
                rel = bron.part.rels[rid]
                blip.set(R_EMBED, nieuw.part.relate_to(rel.target_part, rel.reltype))
        return el

    achtergrond = bron._element.cSld.bg  # een eigen achtergrond (kleur of afbeelding) staat niet in de vormen
    if achtergrond is not None:
        oud = nieuw._element.cSld.bg
        if oud is not None:
            oud.getparent().remove(oud)
        nieuw._element.cSld.insert(0, met_eigen_afbeeldingen(copy.deepcopy(achtergrond)))
    for sh in bron.shapes:
        nieuw.shapes._spTree.append(met_eigen_afbeeldingen(copy.deepcopy(sh._element)))
    kopieer_notities(bron, nieuw)
    verplaats_slide(prs, len(prs.slides), naar)
    prs.save(uit)
    print(f"slide {nummer} gekloond naar positie {naar} -> {uit}")


def kopieer_notities(bron, nieuw):
    if bron.has_notes_slide and bron.notes_slide.notes_text_frame is not None:
        ns = nieuw.notes_slide
        if ns.notes_text_frame is None:  # notitiemaster zonder tekstvak: neem de structuur van de bron over
            for sh in list(ns.shapes):
                sh._element.getparent().remove(sh._element)
            for sh in bron.notes_slide.shapes:
                ns.shapes._spTree.append(copy.deepcopy(sh._element))
        ns.notes_text_frame.text = bron.notes_slide.notes_text_frame.text


def importeer(inp, uit, bronbestand, nummer, naar):
    prs = Presentation(inp)
    bprs = Presentation(bronbestand)
    if not 1 <= nummer <= len(bprs.slides):
        sys.exit(f"FOUT: slide {nummer} bestaat niet in {bronbestand} ({len(bprs.slides)} slides)")
    if not 1 <= naar <= len(prs.slides) + 1:
        sys.exit(f"FOUT: positie {naar} bestaat niet (het deck heeft {len(prs.slides)} slides)")
    bron = bprs.slides[nummer - 1]
    layouts = {l.name: l for l in prs.slide_layouts}
    nieuw = prs.slides.add_slide(layouts.get(bron.slide_layout.name, prs.slide_layouts[0]))
    for sh in list(nieuw.shapes):
        sh._element.getparent().remove(sh._element)

    def met_eigen_relaties(el):
        # afbeeldingen en links verwijzen naar relaties van het bronbestand; leg ze opnieuw aan in dit bestand
        for node in el.iter():
            for attr, rid in list(node.attrib.items()):
                if not attr.startswith(R_NS) or rid not in bron.part.rels:
                    continue
                rel = bron.part.rels[rid]
                if rel.is_external:
                    node.set(attr, nieuw.part.relate_to(rel.target_ref, rel.reltype, is_external=True))
                elif rel.target_part.content_type.startswith("image/"):
                    _, nieuw_rid = nieuw.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                    node.set(attr, nieuw_rid)
                else:
                    sys.exit(f"FOUT: {rel.target_part.partname} ({rel.target_part.content_type}) kan niet mee; "
                             "importeer alleen slides met tekst, vormen, afbeeldingen en tabellen")
        return el

    achtergrond = bron._element.cSld.bg
    if achtergrond is not None:
        nieuw._element.cSld.insert(0, met_eigen_relaties(copy.deepcopy(achtergrond)))
    for sh in bron.shapes:
        nieuw.shapes._spTree.append(met_eigen_relaties(copy.deepcopy(sh._element)))
    kopieer_notities(bron, nieuw)
    verplaats_slide(prs, len(prs.slides), naar)
    prs.save(uit)
    print(f"slide {nummer} uit {bronbestand} geimporteerd op positie {naar} -> {uit}")


def verplaats(inp, uit, nummer, naar):
    prs = Presentation(inp)
    verplaats_slide(prs, nummer, naar)
    prs.save(uit)
    print(f"slide {nummer} verplaatst naar positie {naar} -> {uit}")


def eerste_tekst(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().splitlines()[0][:60]
    return "(geen tekst)"


def verwijder(inp, uit, nummers):
    prs = Presentation(inp)
    totaal = len(prs.slides)
    nummers = sorted(set(nummers), reverse=True)
    fout = [n for n in nummers if not 1 <= n <= totaal]
    if fout:
        sys.exit(f"FOUT: slide {', '.join(map(str, fout))} bestaat niet (het deck heeft {totaal} slides)")
    if len(nummers) == totaal:
        sys.exit("FOUT: je verwijdert dan alle slides")
    lijst = prs.slides._sldIdLst
    for n in nummers:  # van achter naar voren, zodat de nummers van de rest niet verschuiven
        el = list(lijst)[n - 1]
        print(f"  slide {n}: {eerste_tekst(prs.slides[n - 1])}")
        prs.part.drop_rel(el.rId)  # slide en notities hangen nergens meer aan en worden niet meer opgeslagen
        lijst.remove(el)
    prs.save(uit)
    print(f"{len(nummers)} slide(s) verwijderd, {len(prs.slides)} over -> {uit}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("dump"); d.add_argument("inp"); d.add_argument("--notes", action="store_true")
    r = sub.add_parser("replace"); r.add_argument("inp"); r.add_argument("uit"); r.add_argument("map"); r.add_argument("--lenient", action="store_true"); r.add_argument("--slide", type=int, nargs="+")
    b = sub.add_parser("beeldmerk"); b.add_argument("inp"); b.add_argument("uit")
    m = sub.add_parser("media"); m.add_argument("inp"); m.add_argument("uit"); m.add_argument("paren", nargs="+"); m.add_argument("--slide", type=int, nargs="+")
    for naam in ("kloon", "verplaats"):
        k = sub.add_parser(naam); k.add_argument("inp"); k.add_argument("uit")
        k.add_argument("--slide", type=int, required=True); k.add_argument("--naar", type=int, required=True)
    im = sub.add_parser("importeer"); im.add_argument("inp"); im.add_argument("uit"); im.add_argument("bron")
    im.add_argument("--slide", type=int, required=True); im.add_argument("--naar", type=int, required=True)
    v = sub.add_parser("verwijder"); v.add_argument("inp"); v.add_argument("uit")
    v.add_argument("--slide", type=int, nargs="+", required=True)
    a = p.parse_args()
    if a.cmd == "dump":
        dump(a.inp, a.notes)
    elif a.cmd == "replace":
        replace(a.inp, a.uit, a.map, a.lenient, a.slide)
    elif a.cmd == "beeldmerk":
        beeldmerk(a.inp, a.uit)
    elif a.cmd == "media":
        media(a.inp, a.uit, a.paren, a.slide)
    elif a.cmd == "kloon":
        kloon(a.inp, a.uit, a.slide, a.naar)
    elif a.cmd == "verplaats":
        verplaats(a.inp, a.uit, a.slide, a.naar)
    elif a.cmd == "importeer":
        importeer(a.inp, a.uit, a.bron, a.slide, a.naar)
    elif a.cmd == "verwijder":
        verwijder(a.inp, a.uit, a.slide)


if __name__ == "__main__":
    main()
